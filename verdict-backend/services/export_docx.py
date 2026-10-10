from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def build_workspace_docx(workspace: dict[str, Any]) -> bytes:
    document = Document()
    memo = workspace["memo"]
    coverage = workspace.get("coverage_status", {})

    document.add_heading(memo.get("title", "STOCK X Decision Memo"), level=0)
    document.add_paragraph(
        f"Ticker: {memo.get('company_snapshot', {}).get('ticker', workspace.get('ticker', ''))}"
    )
    document.add_paragraph(
        f"Coverage: {coverage.get('market', 'unknown').upper()} | Filings: {coverage.get('filings', 'partial')}"
    )

    snapshot = memo.get("company_snapshot", {})
    document.add_heading("Company Snapshot", level=1)
    document.add_paragraph(snapshot.get("summary", ""))
    _add_bullets(document, snapshot.get("highlights", []))

    sections = [
        ("Core Thesis", memo.get("core_thesis", {})),
        ("Market and Competitive Position", memo.get("market_and_competitive_position", {})),
        ("Financial Quality and Constraints", memo.get("financial_quality_and_constraints", {})),
        ("Key Risks and Stress Points", memo.get("key_risks_and_stress_points", {})),
        ("Recommended Next Moves", memo.get("recommended_next_moves", {})),
    ]

    for title, section in sections:
        document.add_heading(title, level=1)
        document.add_paragraph(section.get("summary", ""))
        _add_bullets(document, section.get("bullets", []))

    qa_history = workspace.get("qa_history") or []
    if qa_history:
        document.add_heading("STOCK X DataPulse Intelligence", level=1)
        for item in qa_history:
            document.add_heading(f"Q: {item.get('question', '')}", level=2)
            document.add_paragraph(item.get("answer", ""))
            takeaways = item.get("key_takeaways", [])
            if takeaways:
                document.add_heading("Key Takeaways", level=3)
                _add_bullets(document, takeaways)
            evidence = item.get("evidence_used", [])
            if evidence:
                document.add_heading("Evidence Cited", level=3)
                for ev in evidence:
                    title_text = ev.get("title", "Document")
                    reason_text = ev.get("reason", "")
                    url_text = ev.get("url", "")
                    p = document.add_paragraph(f"{title_text}: {reason_text}")
                    if url_text:
                        p.add_run(f" ({url_text})").italic = True

    document.add_heading("Evidence and Sources", level=1)
    for source in memo.get("evidence_and_sources", []):
        paragraph = document.add_paragraph()
        paragraph.add_run(source.get("title", "Untitled")).bold = True
        paragraph.add_run(f"\n{source.get('relevance', '')}")
        if source.get("url"):
            paragraph.add_run(f"\n{source['url']}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
